"""
Извлечение текста из загруженных пользователем файлов и подготовка чанков для Pinecone.

Поддерживаемые форматы: .txt, .md, .pdf, .docx.
Старый бинарный .doc не поддерживается — бот попросит пересохранить в .docx.
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass

# --- Настройки чанкинга (можно переопределить через .env) ---
DOC_CHUNK_SIZE: int = int(os.getenv("DOC_CHUNK_SIZE", "1200"))  # целевой размер чанка в символах
DOC_CHUNK_OVERLAP: int = int(os.getenv("DOC_CHUNK_OVERLAP", "200"))  # перекрытие соседних чанков
DOC_MAX_CHUNKS: int = int(os.getenv("DOC_MAX_CHUNKS", "150"))  # верхний предел чанков на один файл
DOC_FILE_SIZE_LIMIT_MB: int = 20  # Telegram Bot API сам отдаёт максимум 20 МБ

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class DocumentContent:
    """Результат извлечения текста из файла."""

    text: str
    extension: str


def normalize_ext(filename: str) -> str:
    """Расширение файла в нижнем регистре с точкой (.pdf, .md и т.п.)."""
    name = filename.strip().lower()
    dot = name.rfind(".")
    if dot <= -1 or dot == len(name) - 1:
        return ""
    return name[dot:]


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        if txt.strip():
            pages.append(txt.strip())
    return "\n\n".join(pages)


def extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = "\t".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def extract_text(data: bytes, filename: str) -> str:
    """Извлечение текста по расширению файла. Пустой результат — если текста нет."""
    ext = normalize_ext(filename)
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext in (".txt", ".md"):
        return decode_text(data)
    raise ValueError(f"Неподдерживаемый формат: {ext or filename}")


def decode_text(data: bytes) -> str:
    """Декодирование текстового файла: utf-8 → cp1251 → latin-1 (никогда не падает)."""
    for encoding in ("utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _split_at_boundary(window: str, text_len_from_start: int) -> int | None:
    """Ищем ближайшую границу (абзац/предложение/перенос) во второй половине окна."""
    for sep in ("\n\n", ". ", "\n"):
        cut = window.rfind(sep)
        if cut > len(window) // 2:
            return text_len_from_start - len(window) + cut + len(sep)
    return None


def chunk_text(
    text: str,
    chunk_size: int = DOC_CHUNK_SIZE,
    overlap: int = DOC_CHUNK_OVERLAP,
) -> list[str]:
    """Разбивка текста на чанки с перекрытием, по возможности по границе абзаца/предложения."""
    clean = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not clean:
        return []

    chunks: list[str] = []
    start = 0
    n = len(clean)
    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            cut = _split_at_boundary(clean[start:end], end)
            if cut is not None and cut > start:
                end = cut
        chunk = clean[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(end - overlap, start + 1)

    return [c for c in chunks if c][:DOC_MAX_CHUNKS]